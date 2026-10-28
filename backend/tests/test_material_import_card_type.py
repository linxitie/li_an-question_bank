"""D1 回归：资料导入的 type 必须忠于输入契约。

缺陷背景（D1，2026-10-27 同伴报告）：xlsx 管理端导入走「正面/背面」双列，
`build_cards()` 已经唯一确定这是卡片，但 confirm 仍对每张卡再跑一次启发式
分类器，卡片正文短 → 判成 article → 卡片落进「备考指南」，翻卡 UI 永不触发。

这里锁死两条行为：
1. xlsx 卡片分支入库 type 恒为 card（分类器不得覆盖确定性事实）；
2. docx 文档分支仍由分类器判定（长文档 → article），不得被本次修复带偏。
"""
import io

import pytest

from app.models.models import StudyMaterial


def _xlsx_bytes(rows):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _docx_bytes(paragraphs):
    from docx import Document
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _upload(client, headers, filename, content):
    return client.post(
        "/api/materials/admin/import/upload",
        files={"file": (filename, content)},
        headers=headers,
    )


class TestXlsxCardImportKeepsCardType:
    """xlsx 正/背面双列即卡片，type 不得被分类器改写。"""

    def test_confirm_creates_card_type_materials(self, client, admin_headers,
                                                 db_session):
        xlsx = _xlsx_bytes([
            ["正面", "背面", "分类"],
            ["逻辑填空三步排", "先判语境再辨词义", "方法技巧"],
            ["资料分析速算", "先看选项再算差值", "方法技巧"],
        ])
        resp = _upload(client, admin_headers, "资料卡片样例.xlsx", xlsx)
        assert resp.status_code == 200, resp.text
        body = resp.json()
        assert body["parse_status"] == "success", body

        resp = client.post(
            f"/api/materials/admin/import/{body['import_id']}/confirm",
            json={
                "subject": "xingce",
                "category": "方法技巧",
                "column_mapping": {"front": 0, "back": 1, "category": 2},
            },
            headers=admin_headers,
        )
        assert resp.status_code == 200, resp.text
        ids = resp.json()["material_ids"]
        assert len(ids) == 2, resp.text

        materials = db_session.query(StudyMaterial).filter(
            StudyMaterial.material_id.in_(ids)).all()
        assert {m.type for m in materials} == {"card"}
        for material in materials:
            assert material.card_front
            assert material.card_back


class TestDocxImportStillClassified:
    """docx 分支仍由分类器判定：长文档应为 article（不回归）。"""

    def test_confirm_long_docx_stays_article(self, client, admin_headers,
                                             db_session):
        # 超过 DOC_SHORT_THRESHOLD(500 字) 才会被判为 article
        docx = _docx_bytes(["行测备考指南"] + ["申论与行测的备考方法说明。" * 60])
        resp = _upload(client, admin_headers, "行测备考指南.docx", docx)
        assert resp.status_code == 200, resp.text
        body = resp.json()
        assert body["parse_status"] == "success", body

        resp = client.post(
            f"/api/materials/admin/import/{body['import_id']}/confirm",
            json={"subject": "xingce", "type": "article"},
            headers=admin_headers,
        )
        assert resp.status_code == 200, resp.text
        ids = resp.json()["material_ids"]
        assert len(ids) == 1, resp.text

        material = db_session.query(StudyMaterial).filter(
            StudyMaterial.material_id == ids[0]).one()
        assert material.type == "article"
